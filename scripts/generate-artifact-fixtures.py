#!/usr/bin/env -S uv run --script
# SPDX-License-Identifier: Apache-2.0
# /// script
# requires-python = ">=3.11"
# dependencies = ["python-docx", "python-pptx", "openpyxl", "reportlab", "pillow", "pypdf"]
# ///
"""Generate real, reopenable document revisions for the artifact storage benchmark.

Run with uv run --script scripts/generate-artifact-fixtures.py --output DIR.
The default fixtures contain one shared 12 MiB incompressible image per document.
They model media-heavy artifacts, not arbitrary office documents or production data.
"""

import argparse
import importlib.metadata
import json
import random
import shutil
from pathlib import Path

from docx import Document
from docx.shared import Inches as DocInches
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as SheetImage
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader
from reportlab.pdfgen import canvas


def generate(folder: Path, image: Path, label: str) -> None:
    folder.mkdir()
    doc = Document()
    doc.add_heading(label, 0)
    doc.add_paragraph("Quarterly analysis: the underlying source image is retained across edits.")
    doc.add_picture(str(image), width=DocInches(5))
    doc.save(folder / "report.docx")

    sheet = Workbook()
    ws = sheet.active
    ws["A1"] = label
    for row in range(2, 1002):
        ws.append([row, row * 1.5, f"=A{row}*B{row}"])
    illustration = SheetImage(str(image))
    illustration.width = illustration.height = 512
    ws.add_image(illustration, "E1")
    sheet.save(folder / "analysis.xlsx")
    sheet.close()

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = label
    slide.shapes.add_picture(str(image), Inches(1), Inches(1.5), width=Inches(5))
    deck.save(folder / "presentation.pptx")

    pdf = canvas.Canvas(str(folder / "report.pdf"), invariant=1)
    pdf.drawString(50, 780, label)
    pdf.drawImage(str(image), 50, 150, width=500, height=500)
    pdf.save()

    # Confirm each generator produced a document its reader can actually open.
    assert Document(folder / "report.docx").paragraphs[0].text == label
    restored = load_workbook(folder / "analysis.xlsx", read_only=True)
    assert restored.active["A1"].value == label
    restored.close()
    assert Presentation(folder / "presentation.pptx").slides[0].shapes.title.text == label
    assert label in PdfReader(folder / "report.pdf").pages[0].extract_text()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--image-size", type=int, default=2048)
    args = parser.parse_args()
    if args.output.exists():
        parser.error("output must be a new directory")
    if not 64 <= args.image_size <= 4096:
        parser.error("image-size must be between 64 and 4096")
    args.output.mkdir(parents=True)
    image = args.output / "source.png"
    for seed, revision, label in [(10, "00-base", "Quarterly report: draft A"),
                                   (10, "02-edit", "Quarterly report: draft B"),
                                   (11, "03-regenerated", "Quarterly report: draft C")]:
        pixels = random.Random(seed).randbytes(args.image_size * args.image_size * 3)
        Image.frombytes("RGB", (args.image_size, args.image_size), pixels).save(image)
        generate(args.output / revision, image, label)
    shutil.copytree(args.output / "00-base", args.output / "01-unchanged")
    image.unlink()
    metadata = {
        "description": "Media-heavy documents; unchanged bytes, text edit with unchanged image, and full image replacement.",
        "image_size": args.image_size,
        "libraries": {name: importlib.metadata.version(name) for name in
                      ["python-docx", "python-pptx", "openpyxl", "reportlab", "pillow", "pypdf"]},
        "file_sizes": {folder.name: {file.name: file.stat().st_size for file in sorted(folder.iterdir())}
                       for folder in sorted(args.output.iterdir()) if folder.is_dir()},
    }
    (args.output / "fixture.json").write_text(json.dumps(metadata, indent=2) + "\n")
    print(json.dumps(metadata, indent=2))


if __name__ == "__main__":
    main()

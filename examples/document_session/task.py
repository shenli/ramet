"""Trusted document task; executes inside the example's pinned image.

SPDX-License-Identifier: Apache-2.0
"""
import io
from pathlib import Path
import sys

sys.path.insert(0, "/workspace/.deps")
import docx
import lxml
from docx import Document


def main():
    if sys.argv[1] == "verify":
        for module in (docx, lxml):
            assert Path(module.__file__).is_relative_to("/workspace/.deps")
        assert docx.__version__ == "1.2.0"
        buffer = io.BytesIO()
        document = Document()
        document.add_paragraph("dependency check")
        document.save(buffer)
        buffer.seek(0)
        assert Document(buffer).paragraphs[0].text == "dependency check"
    elif sys.argv[1] == "start":
        Path("reports").mkdir(exist_ok=True)
        document = Document()
        document.add_heading("Customer report", 0)
        document.add_paragraph("Work saved before sandbox replacement.")
        document.save("reports/report.docx")
    elif sys.argv[1] == "continue":
        document = Document("reports/report.docx")
        assert document.paragraphs[1].text == "Work saved before sandbox replacement."
        document.add_paragraph("Work continued in a new sandbox.")
        document.save("reports/continued.docx")
        assert Document("reports/continued.docx").paragraphs[-1].text == "Work continued in a new sandbox."
    else:
        raise ValueError("expected verify, start, or continue")


if __name__ == "__main__":
    main()
